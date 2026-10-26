"""Rahaf's KPI scorecard (built 2026-10-26).

Daily lines go at the end of the 7:15 PM audit; the weekly section goes in
the Sunday top-up report (Slack only). Everything here reads the per-record
day models that pm_text_sender.load_day() already builds, so Rahaf does no
extra reporting.

Targets stay at the floors until two clean weeks give a baseline. A missed
floor is shown as a gap, never lowered. Conversations and qualified leads
are coached numbers, not effort failures.
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

DIAL_FLOOR = 150
CONV_FLOOR, CONV_GOAL = 10, 20          # calls >= 60s per day
QUAL_FLOOR, QUAL_GOAL = 1, 5            # qualified leads per week (proxy)
CONV_RATE_GOAL = 0.12                   # weekly, rising to 0.20
CONV_RATE_MIN_CN = 100                  # judge the rate only past this many CN calls
BLOCKS = ("10:15", "14:00", "16:45")    # dialing blocks, Central
BLOCK_SLACK_MIN = 45
CORRECT = {"CORRECT", "CORRECT_DNC"}

# Seller-advanced tiers, highest first. Status slugs as used in the account
# (legacy_archive.ACTIVE); a record counts once, at its highest tier reached.
TIERS = (
    ("Offer made / under contract", {"underwrite", "push_deal", "under_contract", "transactions"}),
    ("Hot / qualified", {"Hot Lead", "Warm Lead"}),
    ("Callback scheduled", {"follow_up"}),
)
COACH_DIR = Path(__file__).resolve().parent.parent / "output" / "call_coaching"
COACH_TYPES = ("cold_call", "lead_management", "closer")


def _mark_at_end(m: dict, number: str) -> str | None:
    marks = [x for x in m["marks"] if x["phone"] == number]
    if marks:
        return marks[-1]["to"]
    return (m["phones"].get(number) or {}).get("status")


def _dials(models: list[dict], caller: str) -> list[tuple[dict, dict]]:
    return [(m, c) for m in models for c in m["calls"] if c["by"] == caller and c["dir"] == "outbound"]


def daily_kpis(models: list[dict], caller: str) -> dict:
    dials = _dials(models, caller)
    cn = [(m, c) for m, c in dials if c["result"] == "answered" and _mark_at_end(m, c["to"]) in CORRECT]
    blocks = {}
    for b in BLOCKS:
        h, mi = map(int, b.split(":"))
        centre = h * 60 + mi
        blocks[b] = any(abs(c["ct"].hour * 60 + c["ct"].minute - centre) <= BLOCK_SLACK_MIN for _, c in dials)
    # New-number marking: numbers dialed today that started the day unmarked.
    new, marked = 0, 0
    for m in models:
        for number in {c["to"] for c in m["calls"] if c["by"] == caller and c["dir"] == "outbound"}:
            marks = [x for x in m["marks"] if x["phone"] == number]
            start = marks[0]["frm"] if marks else (m["phones"].get(number) or {}).get("status")
            if start in (None, "", "UNKNOWN"):
                new += 1
                end = marks[-1]["to"] if marks else None
                marked += end not in (None, "", "UNKNOWN")
    return {
        "dials": len(dials), "cn": len(cn),
        "conv60": sum((c["dur"] or 0) >= 60 for _, c in cn), "conv120": sum((c["dur"] or 0) >= 120 for _, c in cn),
        "blocks": blocks, "new_numbers": new, "new_marked": marked,
    }


def _pct(a: int, b: int) -> str:
    return f"{a / b:.0%}" if b else "n/a"


def daily_lines(models: list[dict], caller: str) -> str:
    k = daily_kpis(models, caller)
    gaps = []
    if k["dials"] < DIAL_FLOOR:
        gaps.append(f"dials {k['dials']}/{DIAL_FLOOR}")
    if k["conv60"] < CONV_FLOOR:
        gaps.append(f"conversations {k['conv60']}/{CONV_FLOOR} (coached, not effort)")
    missed = [b for b, hit in k["blocks"].items() if not hit]
    if missed:
        gaps.append("blocks missed " + ", ".join(missed))
    if k["new_numbers"] and k["new_marked"] < k["new_numbers"]:
        gaps.append(f"{k['new_numbers'] - k['new_marked']} new numbers left unmarked")
    out = [
        "", "*KPI scorecard (daily)*",
        f"• Dials {k['dials']} (floor {DIAL_FLOOR}) · CN contacts {k['cn']} · contact rate {_pct(k['cn'], k['dials'])}",
        f"• Conversations ≥60s {k['conv60']} (floor {CONV_FLOOR}, goal {CONV_GOAL}) · ≥120s {k['conv120']} · "
        f"conversation rate {_pct(k['conv60'], k['cn'])}",
        "• Dialing blocks: " + "  ".join(f"{b} {'✓' if hit else '✗'}" for b, hit in k["blocks"].items()),
        f"• New-number marking {k['new_marked']}/{k['new_numbers']} ({_pct(k['new_marked'], k['new_numbers'])})",
        "• Gaps: " + ("; ".join(gaps) if gaps else "none"),
    ]
    return "\n".join(out)


def weekly_kpis(days: dict[str, list[dict]], caller: str) -> dict:
    """`days` maps ISO date -> that day's models (Mon-Fri)."""
    cn = conv60 = 0
    best: dict[str, int] = {}           # record uuid -> best tier index
    tasked, engaged = set(), set()
    photos = offers = 0
    for models in days.values():
        k = daily_kpis(models, caller)
        cn += k["cn"]
        conv60 += k["conv60"]
        for m in models:
            for s in m["status_changes"]:
                for i, (_, slugs) in enumerate(TIERS):
                    if s["to"] in slugs:
                        best[m["uuid"]] = min(best.get(m["uuid"], i), i)
            if any(c["by"] == caller and c["result"] == "answered" and (c["dur"] or 0) >= 60 for c in m["calls"]):
                engaged.add(m["uuid"])
                if m.get("open_tasks"):
                    tasked.add(m["uuid"])
            photos += sum(1 for t, _ in m["tags_added"] if "photo" in t.lower())
            offers += sum(1 for t, _ in m["tags_added"] if t.lower() in ("offer_pending", "offer made"))
    advanced = {label: sum(1 for v in best.values() if v == i) for i, (label, _) in enumerate(TIERS)}
    return {"cn": cn, "conv60": conv60, "advanced": advanced, "engaged": len(engaged),
            "engaged_tasked": len(tasked & engaged), "photos": photos, "offers": offers}


def latest_coaching(kind: str, folder: Path | None = None) -> tuple[str, dict] | None:
    folder = folder or COACH_DIR
    files = sorted(folder.glob(f"rahaf_{kind}_coaching_*.xlsx"))
    if not files:
        return None
    from openpyxl import load_workbook
    ws = load_workbook(files[-1], read_only=True, data_only=True)["Caller Scorecard"]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 2:
        return None
    return files[-1].stem.rsplit("_", 1)[-1], dict(zip(rows[0], rows[1]))


def weekly_lines(days: dict[str, list[dict]], caller: str, today: dt.date | None = None) -> str:
    k = weekly_kpis(days, caller)
    qual = k["advanced"]["Offer made / under contract"] + k["advanced"]["Hot / qualified"]
    rate = (f"{k['conv60'] / k['cn']:.0%} (goal {CONV_RATE_GOAL:.0%})" if k["cn"] >= CONV_RATE_MIN_CN
            else f"not judged yet ({k['cn']}/{CONV_RATE_MIN_CN} CN calls)")
    out = ["", "*KPI scorecard (week)*",
           f"• Conversation rate: {rate}",
           "• Sellers advanced: " + ", ".join(f"{n} {lbl}" for lbl, n in k["advanced"].items()),
           f"• Qualified leads (proxy: hot or better) {qual} (floor {QUAL_FLOOR}, goal {QUAL_GOAL})"
           + ("" if qual >= QUAL_FLOOR else " - gap, coached"),
           f"• Photos received {k['photos']} · offers presented {k['offers']}",
           f"• Engaged sellers with a dated task {k['engaged_tasked']}/{k['engaged']}"]
    for kind in COACH_TYPES:
        got = latest_coaching(kind)
        if got:
            date, row = got
            out.append(f"• Coaching ({kind.replace('_', ' ')}, {date}): avg {row.get('Avg Total')} · "
                       f"weakest {row.get('Recurring Lowest Category')}")
    return "\n".join(out)
